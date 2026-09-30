"""Paso 7. Exporta los datos crudos a un libro de Excel (sin modificar sus valores).

Hojas: Fuentes (manifiesto), Resumen (fórmulas COUNTIFS sobre los datos), SATCAT_pruebas
(filas del catálogo de los fragmentos de las 4 pruebas), GP_Fengyun1C, GP_Kosmos1408,
(sin fórmulas pesadas). La actividad solar va en un libro aparte
(datos_crudos_clima_espacial.xlsx). El catálogo completo (70 mil objetos, usado en el riesgo)
se queda como datos_crudos/satcat.csv por su tamaño. Solo se agregan columnas marcadas como [agregada].
"""
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import comun as c

SALIDA = c.RAIZ / 'datos_crudos_piloto_ASAT.xlsx'
leer = lambda n: pd.read_csv(c.CRUDOS / n, dtype=str, keep_default_na=False)   # texto tal cual

def numerico(df, cols):
    for k in cols:
        df[k] = pd.to_numeric(df[k].replace('', None), errors='coerce')
    return df

sat = numerico(leer('satcat.csv'), ['NORAD_CAT_ID', 'PERIOD', 'INCLINATION', 'APOGEE', 'PERIGEE', 'RCS'])
gp_ch = numerico(leer('gp_fengyun1c_1999-025.csv'), ['MEAN_MOTION', 'ECCENTRICITY', 'INCLINATION', 'RA_OF_ASC_NODE',
                 'ARG_OF_PERICENTER', 'MEAN_ANOMALY', 'NORAD_CAT_ID', 'REV_AT_EPOCH', 'BSTAR', 'MEAN_MOTION_DOT', 'MEAN_MOTION_DDOT'])
gp_ru = numerico(leer('gp_kosmos1408_1982-092.csv'), list(gp_ch.columns[3:9]) + ['NORAD_CAT_ID', 'REV_AT_EPOCH', 'BSTAR', 'MEAN_MOTION_DOT', 'MEAN_MOTION_DDOT'])
sw = leer('SW-All.csv'); sw = numerico(sw, [k for k in sw.columns if k not in ('DATE', 'F10.7_DATA_TYPE')])

# Filas de SATCAT de los fragmentos de cada prueba, con la marca del filtro usado en el análisis
s_full = c.satcat(); filas = []
for p in c.PRUEBAS:
    incl, _, corte = c.fragmentos_de(s_full, p['intdes'], p['fecha'])
    g = sat[sat.OBJECT_ID.str.startswith(p['intdes']) & (sat.OBJECT_TYPE == 'DEB')].copy()
    g.insert(0, 'PRUEBA [agregada]', p['evento'])
    g.insert(1, 'INCLUIDO_EN_ANALISIS [agregada]', g.NORAD_CAT_ID.isin(incl.NORAD_CAT_ID).map({True: 'SI', False: 'NO (anterior a la prueba)'}))
    filas.append(g)
pruebas = pd.concat(filas, ignore_index=True)

man = pd.DataFrame(json.loads((c.CRUDOS / 'manifiesto.json').read_text(encoding='utf-8')))
hojas = {'SATCAT_pruebas': pruebas, 'GP_Fengyun1C': gp_ch, 'GP_Kosmos1408': gp_ru}
with pd.ExcelWriter(SALIDA, engine='openpyxl') as w:
    pd.DataFrame().to_excel(w, sheet_name='Fuentes'); pd.DataFrame().to_excel(w, sheet_name='Resumen')
    for n, df in hojas.items():
        df.to_excel(w, sheet_name=n, index=False)

wb = load_workbook(SALIDA)
AZUL = PatternFill('solid', fgColor='1F4E79'); BLANCA = Font(name='Arial', bold=True, color='FFFFFF', size=10)
NORMAL = Font(name='Arial', size=10); NEGRITA = Font(name='Arial', size=10, bold=True)

# Fuentes
f = wb['Fuentes']; f.delete_rows(1, f.max_row)
f['A1'] = 'Datos crudos del piloto: fragmentos de pruebas antisatélite'; f['A1'].font = Font(name='Arial', bold=True, size=13)
f['A2'] = 'Valores copiados sin modificar de los archivos descargados. Las columnas marcadas [agregada] no están en la fuente original.'
enc = ['Archivo', 'URL', 'Descargado (UTC)', 'Bytes', 'SHA-256', 'Hoja en este libro']
hoja_de = {'satcat.csv': 'SATCAT_pruebas (catálogo completo: datos_crudos/satcat.csv)', 'gp_fengyun1c_1999-025.csv': 'GP_Fengyun1C',
           'gp_kosmos1408_1982-092.csv': 'GP_Kosmos1408', 'SW-All.csv': 'Libro aparte: datos_crudos_clima_espacial.xlsx'}
for j, h in enumerate(enc, 1):
    f.cell(4, j, h)
for i, r in enumerate(man.itertuples(), 5):
    for j, v in enumerate([r.archivo, r.url, r.descargado_utc, r.bytes, r.sha256, hoja_de[r.archivo]], 1):
        f.cell(i, j, v)
notas = [
    ('Columnas principales', ''),
    ('SATCAT: OBJECT_ID', 'Designador internacional (año-lanzamiento-pieza); los fragmentos heredan el del satélite destruido'),
    ('SATCAT: NORAD_CAT_ID', 'Número de catálogo; crece con el tiempo. Se usa para separar fragmentos anteriores y posteriores a la prueba'),
    ('SATCAT: DECAY_DATE', 'Fecha de reingreso; vacía = sigue en órbita'),
    ('SATCAT: PERIOD / APOGEE / PERIGEE', 'Periodo (min) y alturas (km) de la última órbita conocida'),
    ('GP: MEAN_MOTION', 'Movimiento medio n (vueltas/día); con la 3a ley de Kepler da a = (mu/n^2)^(1/3)'),
    ('GP: ECCENTRICITY', 'Excentricidad e de la órbita'),
    ('GP: BSTAR', 'Término de frenado B* (1/radio terrestre) del modelo SGP4; no es un coeficiente físico (ver README)'),
    ('GP: MEAN_MOTION_DOT', 'Mitad de la tasa de cambio del movimiento medio, (dn/dt)/2, en vueltas/día^2'),
    ('Clima: F10.7_OBS / F10.7_OBS_CENTER81', 'Flujo solar a 10.7 cm (diario / media de 81 días), en unidades de flujo solar'),
    ('Clima: AP_AVG', 'Índice geomagnético Ap diario'),
    ('Clima: F10.7_DATA_TYPE', 'OBS = observado; PRD/PRM = predicción diaria/mensual'),
]
for i, (a, b) in enumerate(notas, 5 + len(man) + 2):
    f.cell(i, 1, a); f.cell(i, 2, b)
    f.cell(i, 1).font = NEGRITA if b == '' else NORMAL
for w_, col in zip([34, 62, 24, 11, 68, 34], 'ABCDEF'):
    f.column_dimensions[col].width = w_

# Resumen con fórmulas sobre la hoja SATCAT_pruebas
r = wb['Resumen']; r.delete_rows(1, r.max_row)
r['A1'] = 'Resumen calculado con fórmulas sobre la hoja SATCAT_pruebas'; r['A1'].font = Font(name='Arial', bold=True, size=13)
r['A2'] = 'Solo fragmentos con INCLUIDO_EN_ANALISIS = "SI". "En órbita" = DECAY_DATE vacía.'
enc = ['Prueba', 'Satélite', 'Altura aprox. (km)', 'Fragmentos', 'En órbita hoy', '% en órbita', 'Excluidos (anteriores)']
for j, h in enumerate(enc, 1):
    r.cell(4, j, h)
ult = len(pruebas) + 1
colP, colI = "SATCAT_pruebas!$A$2:$A$%d" % ult, "SATCAT_pruebas!$B$2:$B$%d" % ult
idx_dec = list(pruebas.columns).index('DECAY_DATE') + 1
colD = f"SATCAT_pruebas!${get_column_letter(idx_dec)}$2:${get_column_letter(idx_dec)}${ult}"
for i, p in enumerate(c.PRUEBAS, 5):
    r.cell(i, 1, p['evento']); r.cell(i, 2, p['satelite']); r.cell(i, 3, p['altitud_km'])
    r.cell(i, 4, f'=COUNTIFS({colP},A{i},{colI},"SI")')
    r.cell(i, 5, f'=COUNTIFS({colP},A{i},{colI},"SI",{colD},"")')
    r.cell(i, 6, f'=IF(D{i}=0,0,E{i}/D{i})'); r.cell(i, 6).number_format = '0.0%'
    r.cell(i, 7, f'=COUNTIFS({colP},A{i},{colI},"NO*")')
r.cell(10, 1, 'Nota: la altura aproximada es un valor de referencia de la literatura, no sale de estas hojas.').font = Font(name='Arial', italic=True, size=9)
for w_, col in zip([16, 14, 18, 12, 14, 12, 22], 'ABCDEFG'):
    r.column_dimensions[col].width = w_

# Formato común
for ws in wb.worksheets:
    fila_enc = 4 if ws.title in ('Fuentes', 'Resumen') else 1
    for fila in ws.iter_rows():
        for cel in fila:
            if cel.font.name != 'Arial' or cel.row not in (1,):
                if not (ws.title in ('Fuentes', 'Resumen') and cel.row == 1):
                    cel.font = Font(name='Arial', size=10, bold=cel.font.bold, italic=cel.font.italic)
    for cel in ws[fila_enc]:
        if cel.value is not None:
            cel.font = BLANCA; cel.fill = AZUL; cel.alignment = Alignment(vertical='center', wrap_text=True)
    if ws.title not in ('Fuentes', 'Resumen'):
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
            ancho = max(len(str(x.value)) if x.value is not None else 0 for x in col)
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, ancho + 2), 40)
wb.calculation.fullCalcOnLoad = True   # Excel/Sheets recalculan las fórmulas al abrir
wb.save(SALIDA)

# Libro aparte para la actividad solar (25 mil filas, sin fórmulas)
SALIDA_SW = c.RAIZ / 'datos_crudos_clima_espacial.xlsx'
with pd.ExcelWriter(SALIDA_SW, engine='openpyxl') as w:
    sw.to_excel(w, sheet_name='Clima_espacial', index=False)
wb2 = load_workbook(SALIDA_SW); ws = wb2['Clima_espacial']
for fila in ws.iter_rows():
    for cel in fila:
        cel.font = NORMAL
for cel in ws[1]:
    cel.font = BLANCA; cel.fill = AZUL
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
for j in range(1, ws.max_column + 1):
    ws.column_dimensions[get_column_letter(j)].width = 14
wb2.save(SALIDA_SW)
print('Guardado', SALIDA, {n: len(d) for n, d in hojas.items()}, 'y', SALIDA_SW, len(sw))
